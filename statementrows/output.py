"""Write a read statement out as CSV, Excel and JSON."""
from __future__ import annotations

import csv
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .core import Result

COLUMNS = ["date", "description", "amount", "direction", "balance", "check", "account", "page", "date_2", "note", "ref"]


def rows(res: Result) -> list[dict]:
    return [{
        "date": t.date.isoformat() if t.date else "",
        "description": t.description,
        "amount": f"{t.amount:.2f}",
        "direction": t.direction,
        "balance": f"{t.balance:.2f}" if t.balance is not None else "",
        "check": t.check,
        "account": t.account,
        "page": t.page,
        "date_2": t.date_2.isoformat() if t.date_2 else "",
        "note": t.note,
        "ref": t.ref,
    } for t in res.txns]


def summary(res: Result) -> dict:
    money_in = sum((t.amount for t in res.txns if t.amount > 0), Decimal(0))
    money_out = sum((t.amount for t in res.txns if t.amount < 0), Decimal(0))
    return {
        "file": Path(res.path).name,
        "status": res.status,
        "kind": res.kind,
        "rows": len(res.txns),
        "opening": f"{res.opening:.2f}" if res.opening is not None else "",
        "closing": f"{res.closing:.2f}" if res.closing is not None else "",
        "money_in": f"{money_in:.2f}",
        "money_out": f"{money_out:.2f}",
        "period": f"{res.period[0]} to {res.period[1]}" if res.period else "",
        "checks": res.checks,
        "problems": res.problems,
    }


def write_csv(res: Result, path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows(res))


def write_json(res: Result, path: Path) -> None:
    path.write_text(json.dumps({"summary": summary(res), "transactions": rows(res)}, indent=2,
                               ensure_ascii=False), encoding="utf-8")


def write_xlsx(res: Result, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    bad = PatternFill("solid", fgColor="F8D7DA")
    for t in res.txns:
        ws.append([t.date, t.description, float(t.amount), t.direction,
                   float(t.balance) if t.balance is not None else None, t.check, t.account, t.page, t.date_2,
                   t.note, t.ref])
        r = ws.max_row
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 9).number_format = "yyyy-mm-dd"
        ws.cell(r, 3).number_format = "#,##0.00"
        ws.cell(r, 5).number_format = "#,##0.00"
        if t.check == "break" or t.date is None:
            for c in ws[r]:
                c.fill = bad
    widths = {"A": 12, "B": 60, "C": 13, "D": 10, "E": 14, "F": 24, "G": 30, "H": 6, "I": 12, "J": 20, "K": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    s = wb.create_sheet("Check")
    sm = summary(res)
    for k in ("file", "status", "kind", "rows", "opening", "closing", "money_in", "money_out", "period"):
        s.append([k, sm[k]])
    s.append([])
    s.append(["what was checked"])
    s.cell(s.max_row, 1).font = Font(bold=True)
    for c in res.checks:
        s.append(["ok", c])
    s.append([])
    s.append(["what did not add up"])
    s.cell(s.max_row, 1).font = Font(bold=True)
    for p in res.problems or ["nothing"]:
        s.append(["!" if res.problems else "", p])
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 110
    wb.save(path)


def write_all(res: Result, out_dir: Path, stem: str) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = [out_dir / f"{stem}.xlsx", out_dir / f"{stem}.csv", out_dir / f"{stem}.json"]
    write_xlsx(res, paths[0])
    write_csv(res, paths[1])
    write_json(res, paths[2])
    return paths
